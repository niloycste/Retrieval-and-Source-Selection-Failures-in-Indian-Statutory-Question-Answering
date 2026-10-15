"""Add Gwet's AC1 and weighted materiality kappa to the audit workbook.

The primary reliability analysis uses the original, pre-consensus labels in
``Frozen_Input_A`` and ``Original_Input_C``.  The script also calculates the
same descriptive statistics for the current revised snapshots in
``Frozen_Input_A`` and ``Frozen_Input_C``.

Before writing anything, the script reconstructs the normalized
annotated/selected roles and verifies that raw agreement and nominal Cohen's
kappa reproduce every existing row in Agreement_Original and
Agreement_Frozen.  It never modifies either annotator workbook, the frozen
inputs, or the consensus decisions.
"""

from __future__ import annotations

import argparse
import math
import shutil
from collections import Counter
from copy import copy
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = ROOT / "outputs" / "human validation" / "legal_expert_audit_master.xlsx"

ROLE_FIELDS = {
    "Question uniquely identifies annotated source": (
        "unique",
        "annotated",
        ["Yes", "No", "Unclear"],
    ),
    "Question uniquely identifies selected source": (
        "unique",
        "selected",
        ["Yes", "No", "Unclear"],
    ),
    "Annotated source legally relevant": (
        "relevance",
        "annotated",
        ["Yes", "Partly", "No", "Unclear"],
    ),
    "Selected source legally relevant": (
        "relevance",
        "selected",
        ["Yes", "Partly", "No", "Unclear"],
    ),
    "Annotated answer legally correct": (
        "correctness",
        "annotated",
        ["Yes", "Partly", "No", "Unclear"],
    ),
    "Predicted answer legally correct": (
        "correctness",
        "selected",
        ["Yes", "Partly", "No", "Unclear"],
    ),
    "Predicted-answer error materiality": (
        "materiality",
        "selected",
        ["None", "Low", "Moderate", "High", "Not applicable", "Unclear"],
    ),
    "Annotated-answer error materiality": (
        "materiality",
        "annotated",
        ["None", "Low", "Moderate", "High", "Not applicable", "Unclear"],
    ),
}

CASE_FIELDS = {
    "Sources legally equivalent": (
        "Passages legally equivalent",
        ["Yes", "Partly", "No", "Not applicable", "Unclear"],
    ),
    "Text quality issue": (
        "Text quality issue",
        ["None", "Minor", "Material", "Unclear"],
    ),
}

DISPLAY_FIELDS = {
    "unique": {
        "A": "Passage A uniquely identified by question",
        "B": "Passage B uniquely identified by question",
    },
    "relevance": {
        "A": "Passage A legally relevant",
        "B": "Passage B legally relevant",
    },
    "correctness": {
        "A": "Answer A legally correct",
        "B": "Answer B legally correct",
    },
    "materiality": {
        "A": "Answer A error materiality",
        "B": "Answer B error materiality",
    },
}

MATERIALITY_ORDER = ["None", "Low", "Moderate", "High"]


def sheet_records(sheet) -> list[dict]:
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value or "").strip() for value in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:]]


def text(value) -> str:
    return str(value or "").strip()


def cohen_kappa(left: list[str], right: list[str], labels: list[str]) -> float | None:
    n = len(left)
    if not n:
        return None
    observed = sum(a == b for a, b in zip(left, right)) / n
    left_counts, right_counts = Counter(left), Counter(right)
    expected = sum(left_counts[label] * right_counts[label] for label in labels) / (n * n)
    return None if math.isclose(expected, 1.0) else (observed - expected) / (1 - expected)


def gwet_ac1(left: list[str], right: list[str], labels: list[str]) -> float | None:
    """Gwet's AC1 for two raters and nominal categories."""
    n, categories = len(left), len(labels)
    if not n or categories < 2:
        return None
    observed = sum(a == b for a, b in zip(left, right)) / n
    left_counts, right_counts = Counter(left), Counter(right)
    pooled = [
        (left_counts[label] + right_counts[label]) / (2 * n)
        for label in labels
    ]
    expected = sum(probability * (1 - probability) for probability in pooled) / (categories - 1)
    return None if math.isclose(expected, 1.0) else (observed - expected) / (1 - expected)


def quadratic_weighted_kappa(
    left: list[str], right: list[str], labels: list[str]
) -> tuple[float | None, int, int]:
    pairs = [(a, b) for a, b in zip(left, right) if a in labels and b in labels]
    excluded = len(left) - len(pairs)
    if not pairs:
        return None, 0, excluded
    n = len(pairs)
    final_index = max(len(labels) - 1, 1)
    index = {label: position for position, label in enumerate(labels)}

    def weight(a: str, b: str) -> float:
        return 1 - ((index[a] - index[b]) / final_index) ** 2

    observed = sum(weight(a, b) for a, b in pairs) / n
    left_counts = Counter(a for a, _ in pairs)
    right_counts = Counter(b for _, b in pairs)
    expected = sum(
        weight(left_label, right_label)
        * left_counts[left_label]
        * right_counts[right_label]
        / (n * n)
        for left_label in labels
        for right_label in labels
    )
    value = None if math.isclose(expected, 1.0) else (observed - expected) / (1 - expected)
    return value, n, excluded


def load_mapping(workbook) -> tuple[dict, dict]:
    metadata = {text(row["case_id"]): row for row in sheet_records(workbook["Metadata_Key"])}
    mapping = {}
    for row in sheet_records(workbook["Import_Mapping"]):
        mapping[(text(row["Case ID"]), text(row["Source code"]))] = text(
            row["Original file Passage A role"]
        )
    return metadata, mapping


def normalized_records(workbook, source_c_sheet: str) -> dict[str, dict]:
    metadata, mapping = load_mapping(workbook)
    source_rows = {
        "A": {text(row["Case ID"]): row for row in sheet_records(workbook["Frozen_Input_A"])},
        "C": {text(row["Case ID"]): row for row in sheet_records(workbook[source_c_sheet])},
    }
    normalized = {}
    for case_id, meta in metadata.items():
        normalized[case_id] = {"stage": text(meta["stage_outcome"]), "A": {}, "C": {}}
        for source_code in ("A", "C"):
            row = source_rows[source_code][case_id]
            passage_a_role = mapping[(case_id, source_code)]
            if passage_a_role == "indistinguishable":
                # In identical-candidate rows, either orientation is valid
                # when the two labels agree.  The few asymmetric historical
                # fields are explicitly excluded below.
                passage_a_role = "annotated"
            passage_b_role = "selected" if passage_a_role == "annotated" else "annotated"
            role_to_display = {passage_a_role: "A", passage_b_role: "B"}
            for family, fields in DISPLAY_FIELDS.items():
                label_a = text(row[fields["A"]])
                label_b = text(row[fields["B"]])
                # The historical audit records eight specific original C
                # judgments whose annotated/selected role could not be
                # recovered without guessing.  Agreement_Original excludes
                # exactly those cells; all other roles use the preserved
                # decoding map even when the displayed evidence is identical.
                original_exclusions = {
                    ("LEX-050", "unique"),
                    ("LEX-053", "materiality"),
                    ("LEX-060", "correctness"),
                    ("LEX-060", "materiality"),
                }
                ambiguous = (
                    source_code == "C"
                    and source_c_sheet == "Original_Input_C"
                    and (case_id, family) in original_exclusions
                )
                for role, display in role_to_display.items():
                    normalized[case_id][source_code][(family, role)] = (
                        None if ambiguous else text(row[fields[display]])
                    )
            for judgment, (field, _) in CASE_FIELDS.items():
                normalized[case_id][source_code][("case", judgment)] = text(row[field])
    return normalized


def values_for(
    normalized: dict[str, dict], scope: str, judgment: str
) -> tuple[list[str], list[str], list[str]]:
    if judgment in ROLE_FIELDS:
        family, role, labels = ROLE_FIELDS[judgment]
        key = (family, role)
    else:
        _, labels = CASE_FIELDS[judgment]
        key = ("case", judgment)
    left, right = [], []
    for record in normalized.values():
        if scope != "All 80 supplied cases" and record["stage"] != scope:
            continue
        a, c = record["A"].get(key), record["C"].get(key)
        if a is None or c is None:
            continue
        if a not in labels or c not in labels:
            raise ValueError(f"Invalid labels for {judgment}: {a!r}, {c!r}")
        left.append(a)
        right.append(c)
    return left, right, labels


def close(left: float | None, right: float | None, tolerance: float = 1e-10) -> bool:
    if left is None or right is None:
        return left is None and right is None
    return math.isclose(left, right, rel_tol=tolerance, abs_tol=tolerance)


def update_agreement_sheet(sheet, normalized: dict[str, dict]) -> None:
    headers = [cell.value for cell in sheet[1]]
    expected = [
        "Scope / stratum", "Judgment", "Paired cases", "Exact agreements",
        "Disagreements", "Raw agreement", "Expected agreement",
        "Cohen kappa (nominal)", "Interpretation / calculation basis",
    ]
    if headers[:9] != expected:
        raise ValueError(f"Unexpected header layout in {sheet.title}: {headers[:9]}")

    new_headers = [
        "Gwet AC1",
        "Quadratic weighted kappa",
        "Weighted-kappa paired cases",
        "Excluded from weighted kappa",
    ]
    for offset, value in enumerate(new_headers, start=10):
        target = sheet.cell(1, offset)
        target.value = value
        source = sheet.cell(1, 9)
        target._style = copy(source._style)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)

    for row_number in range(2, sheet.max_row + 1):
        scope = text(sheet.cell(row_number, 1).value)
        judgment = text(sheet.cell(row_number, 2).value)
        left, right, labels = values_for(normalized, scope, judgment)
        n = len(left)
        exact = sum(a == b for a, b in zip(left, right))
        raw = exact / n if n else None
        nominal = cohen_kappa(left, right, labels)

        # Reproduce the existing results before adding new metrics.
        if sheet.cell(row_number, 3).value != n:
            raise ValueError(f"{sheet.title} row {row_number}: paired-case mismatch")
        if sheet.cell(row_number, 4).value != exact:
            raise ValueError(f"{sheet.title} row {row_number}: exact-agreement mismatch")
        if not close(sheet.cell(row_number, 6).value, raw):
            raise ValueError(f"{sheet.title} row {row_number}: raw-agreement mismatch")
        if not close(sheet.cell(row_number, 8).value, nominal):
            raise ValueError(f"{sheet.title} row {row_number}: Cohen-kappa mismatch")

        ac1 = gwet_ac1(left, right, labels)
        sheet.cell(row_number, 10).value = ac1
        if ac1 is not None:
            sheet.cell(row_number, 10).number_format = "0.000000"

        if "materiality" in judgment.lower():
            weighted, weighted_n, excluded = quadratic_weighted_kappa(
                left, right, MATERIALITY_ORDER
            )
            sheet.cell(row_number, 11).value = weighted
            sheet.cell(row_number, 12).value = weighted_n
            sheet.cell(row_number, 13).value = excluded
            if weighted is not None:
                sheet.cell(row_number, 11).number_format = "0.000000"
        else:
            sheet.cell(row_number, 11).value = None
            sheet.cell(row_number, 12).value = None
            sheet.cell(row_number, 13).value = None

        for column in range(10, 14):
            source = sheet.cell(row_number, 8 if column < 13 else 9)
            target = sheet.cell(row_number, column)
            target._style = copy(source._style)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)

    for column in range(10, 14):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 24
    sheet.auto_filter.ref = f"A1:M{sheet.max_row}"


def add_revision_log_entry(workbook) -> None:
    sheet = workbook["Revision_Log"]
    existing = {
        (text(row[0]), text(row[1]), text(row[2]))
        for row in sheet.iter_rows(min_row=2, values_only=True)
    }
    key = ("Agreement metrics", "Agreement_Original and Agreement_Frozen", "Gwet AC1 and weighted kappa")
    if key in existing:
        return
    row = sheet.max_row + 1
    values = [
        key[0], key[1], key[2],
        "Nominal Cohen kappa only",
        "Gwet AC1 added for all categorical judgments; quadratic weighted kappa added for ordered materiality",
        "Calculated reproducibly from the same paired labels and exclusions used by each agreement sheet.",
    ]
    for column, value in enumerate(values, start=1):
        sheet.cell(row, column).value = value
        sheet.cell(row, column)._style = copy(sheet.cell(row - 1, column)._style)
        sheet.cell(row, column).alignment = copy(sheet.cell(row - 1, column).alignment)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--no-backup", action="store_true")
    args = parser.parse_args()
    workbook_path = args.workbook.resolve()
    if not workbook_path.exists():
        raise FileNotFoundError(workbook_path)

    backup = workbook_path.with_name(workbook_path.stem + ".before_agreement_metrics.xlsx")
    if not args.no_backup and not backup.exists():
        shutil.copy2(workbook_path, backup)

    workbook = openpyxl.load_workbook(workbook_path)
    original = normalized_records(workbook, "Original_Input_C")
    current = normalized_records(workbook, "Frozen_Input_C")
    update_agreement_sheet(workbook["Agreement_Original"], original)
    update_agreement_sheet(workbook["Agreement_Frozen"], current)
    add_revision_log_entry(workbook)
    workbook.save(workbook_path)

    print(f"Updated: {workbook_path}")
    print(f"Backup:  {backup if backup.exists() else 'disabled'}")
    print("Verified existing raw agreement and Cohen kappa before writing AC1/weighted kappa.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
